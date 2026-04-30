"""Report generation and export functionality.

Provides PDF and Excel export for compliance reports, violation summaries,
training progress, and analytics.
"""
import logging
from typing import Dict, Any, List, Optional, cast
from datetime import datetime, timedelta
from io import BytesIO
from django.http import HttpResponse
from django.template.loader import render_to_string
from django.contrib.auth import get_user_model
from django.db.models import Count, Q, Avg
from django.utils import timezone

logger = logging.getLogger(__name__)

User = get_user_model()

HTML = None
CSS = None
FontConfiguration = None

try:
    from weasyprint import HTML, CSS
    from weasyprint.text.fonts import FontConfiguration
    WEASYPRINT_AVAILABLE = True
except ImportError:
    WEASYPRINT_AVAILABLE = False
    logger.warning('WeasyPrint not installed - PDF export disabled')

openpyxl = None
Font = Alignment = PatternFill = Border = Side = None
BarChart = PieChart = Reference = None

try:
    import openpyxl
    from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
    from openpyxl.chart import BarChart, PieChart, Reference
    OPENPYXL_AVAILABLE = True
except ImportError:
    OPENPYXL_AVAILABLE = False
    logger.warning('openpyxl not installed - Excel export disabled')


def _font(*args, **kwargs):
    if Font is None:
        raise ImportError('openpyxl required for Excel styling')
    return cast(Any, Font)(*args, **kwargs)


def _pattern_fill(*args, **kwargs):
    if PatternFill is None:
        raise ImportError('openpyxl required for Excel styling')
    return cast(Any, PatternFill)(*args, **kwargs)


class ComplianceReportGenerator:
    """Generate compliance reports in PDF and Excel formats."""
    
    def __init__(self, start_date: Optional[datetime] = None, end_date: Optional[datetime] = None):
        """
        Initialize report generator.
        
        Args:
            start_date: Report start date (default: 30 days ago)
            end_date: Report end date (default: now)
        """
        self.end_date = end_date or timezone.now()
        self.start_date = start_date or (self.end_date - timedelta(days=30))
    
    def get_violation_data(self) -> Dict[str, Any]:
        """Get violation statistics for the report period."""
        from policy.models import Violation
        
        violations = Violation.objects.filter(
            detected_at__gte=self.start_date,
            detected_at__lte=self.end_date
        )
        
        return {
            'total': violations.count(),
            'by_severity': {
                'critical': violations.filter(severity='critical').count(),
                'high': violations.filter(severity='high').count(),
                'medium': violations.filter(severity='medium').count(),
                'low': violations.filter(severity='low').count(),
            },
            'by_status': {
                'open': violations.filter(status='open').count(),
                'acknowledged': violations.filter(status='acknowledged').count(),
                'resolved': violations.filter(status='resolved').count(),
            },
            'top_violators': violations.values('user__username').annotate(
                count=Count('id')
            ).order_by('-count')[:10],
            'recent_violations': violations.order_by('-detected_at')[:20],
        }
    
    def get_training_data(self) -> Dict[str, Any]:
        """Get training completion statistics."""
        from training.models import TrainingModule, TrainingProgress
        
        modules = TrainingModule.objects.all()
        total_users = User.objects.filter(is_active=True).count()
        
        module_stats = []
        for module in modules:
            completed = TrainingProgress.objects.filter(
                module=module,
                completed_at__gte=self.start_date,
                completed_at__lte=self.end_date
            ).count()
            
            completion_rate = (completed / total_users * 100) if total_users > 0 else 0
            
            module_stats.append({
                'title': module.title,
                'completed': completed,
                'total_users': total_users,
                'completion_rate': completion_rate,
            })
        
        return {
            'total_modules': modules.count(),
            'module_stats': module_stats,
            'overall_completion': sum(m['completion_rate'] for m in module_stats) / len(module_stats) if module_stats else 0,
        }
    
    def get_quiz_data(self) -> Dict[str, Any]:
        """Get quiz performance statistics."""
        from quizzes.models import Quiz, QuizAttempt
        
        attempts = QuizAttempt.objects.filter(
            started_at__gte=self.start_date,
            started_at__lte=self.end_date,
            completed_at__isnull=False
        )
        
        return {
            'total_attempts': attempts.count(),
            'passed': attempts.filter(passed=True).count(),
            'failed': attempts.filter(passed=False).count(),
            'average_score': attempts.aggregate(Avg('score'))['score__avg'] or 0,
            'pass_rate': (attempts.filter(passed=True).count() / attempts.count() * 100) if attempts.count() > 0 else 0,
        }
    
    def generate_pdf_report(self) -> bytes:
        """Generate PDF compliance report."""
        if not WEASYPRINT_AVAILABLE or HTML is None or FontConfiguration is None:
            raise ImportError('WeasyPrint required for PDF generation')
        
        # Gather data
        data = {
            'report_date': timezone.now(),
            'start_date': self.start_date,
            'end_date': self.end_date,
            'violations': self.get_violation_data(),
            'training': self.get_training_data(),
            'quizzes': self.get_quiz_data(),
        }
        
        # Render HTML template
        html_string = render_to_string('reports/compliance_report.html', data)
        
        # Generate PDF
        font_configuration_cls = cast(Any, FontConfiguration)
        html_cls = cast(Any, HTML)
        font_config = font_configuration_cls()
        html = html_cls(string=html_string)
        pdf_bytes = html.write_pdf(font_config=font_config)
        if pdf_bytes is None:
            raise RuntimeError('PDF generation returned no content')
        
        return pdf_bytes
    
    def generate_excel_report(self) -> bytes:
        """Generate Excel compliance report."""
        if not OPENPYXL_AVAILABLE or openpyxl is None:
            raise ImportError('openpyxl required for Excel generation')
        
        # Create workbook
        workbook_module = cast(Any, openpyxl)
        wb = workbook_module.Workbook()
        active_sheet = wb.active
        if active_sheet is not None:
            wb.remove(active_sheet)
        
        # Gather data
        violations = self.get_violation_data()
        training = self.get_training_data()
        quizzes = self.get_quiz_data()
        
        # Create Summary sheet
        self._create_summary_sheet(wb, violations, training, quizzes)
        
        # Create Violations sheet
        self._create_violations_sheet(wb, violations)
        
        # Create Training sheet
        self._create_training_sheet(wb, training)
        
        # Create Quizzes sheet
        self._create_quizzes_sheet(wb, quizzes)
        
        # Save to BytesIO
        output = BytesIO()
        wb.save(output)
        output.seek(0)
        
        return output.getvalue()
    
    def _create_summary_sheet(self, wb, violations, training, quizzes):
        """Create summary sheet in Excel workbook."""
        ws = wb.create_sheet('Summary', 0)
        
        # Header
        ws['A1'] = 'Compliance Report Summary'
        ws['A1'].font = _font(size=16, bold=True)
        ws['A2'] = f'Period: {self.start_date.strftime("%Y-%m-%d")} to {self.end_date.strftime("%Y-%m-%d")}'
        
        # Violations summary
        ws['A4'] = 'Violations'
        ws['A4'].font = _font(bold=True)
        ws['A5'] = 'Total Violations:'
        ws['B5'] = violations['total']
        ws['A6'] = 'Critical:'
        ws['B6'] = violations['by_severity']['critical']
        ws['A7'] = 'High:'
        ws['B7'] = violations['by_severity']['high']
        ws['A8'] = 'Medium:'
        ws['B8'] = violations['by_severity']['medium']
        ws['A9'] = 'Low:'
        ws['B9'] = violations['by_severity']['low']
        
        # Training summary
        ws['A11'] = 'Training'
        ws['A11'].font = _font(bold=True)
        ws['A12'] = 'Total Modules:'
        ws['B12'] = training['total_modules']
        ws['A13'] = 'Overall Completion:'
        ws['B13'] = f"{training['overall_completion']:.1f}%"
        
        # Quiz summary
        ws['A15'] = 'Quizzes'
        ws['A15'].font = _font(bold=True)
        ws['A16'] = 'Total Attempts:'
        ws['B16'] = quizzes['total_attempts']
        ws['A17'] = 'Pass Rate:'
        ws['B17'] = f"{quizzes['pass_rate']:.1f}%"
        ws['A18'] = 'Average Score:'
        ws['B18'] = f"{quizzes['average_score']:.1f}%"
        
        # Adjust column widths
        ws.column_dimensions['A'].width = 25
        ws.column_dimensions['B'].width = 15
    
    def _create_violations_sheet(self, wb, violations):
        """Create violations sheet in Excel workbook."""
        ws = wb.create_sheet('Violations')
        
        # Headers
        headers = ['Date', 'User', 'Rule', 'Severity', 'Status']
        for col, header in enumerate(headers, 1):
            cell = ws.cell(row=1, column=col, value=header)
            cell.font = _font(bold=True)
            cell.fill = _pattern_fill(start_color='366092', end_color='366092', fill_type='solid')
            cell.font = _font(color='FFFFFF', bold=True)
        
        # Data
        for row, violation in enumerate(violations['recent_violations'], 2):
            ws.cell(row=row, column=1, value=violation.detected_at.strftime('%Y-%m-%d %H:%M'))
            ws.cell(row=row, column=2, value=violation.user.username)
            ws.cell(row=row, column=3, value=violation.rule.description)
            ws.cell(row=row, column=4, value=violation.severity)
            ws.cell(row=row, column=5, value=violation.status)
        
        # Adjust column widths
        for col in range(1, 6):
            ws.column_dimensions[chr(64 + col)].width = 20
    
    def _create_training_sheet(self, wb, training):
        """Create training sheet in Excel workbook."""
        ws = wb.create_sheet('Training')
        
        # Headers
        headers = ['Module', 'Completed', 'Total Users', 'Completion Rate']
        for col, header in enumerate(headers, 1):
            cell = ws.cell(row=1, column=col, value=header)
            cell.font = _font(bold=True)
            cell.fill = _pattern_fill(start_color='366092', end_color='366092', fill_type='solid')
            cell.font = _font(color='FFFFFF', bold=True)
        
        # Data
        for row, module in enumerate(training['module_stats'], 2):
            ws.cell(row=row, column=1, value=module['title'])
            ws.cell(row=row, column=2, value=module['completed'])
            ws.cell(row=row, column=3, value=module['total_users'])
            ws.cell(row=row, column=4, value=f"{module['completion_rate']:.1f}%")
        
        # Adjust column widths
        ws.column_dimensions['A'].width = 40
        for col in range(2, 5):
            ws.column_dimensions[chr(64 + col)].width = 15
    
    def _create_quizzes_sheet(self, wb, quizzes):
        """Create quizzes sheet in Excel workbook."""
        ws = wb.create_sheet('Quizzes')
        
        # Summary data
        ws['A1'] = 'Quiz Performance Summary'
        ws['A1'].font = _font(size=14, bold=True)
        
        ws['A3'] = 'Total Attempts:'
        ws['B3'] = quizzes['total_attempts']
        ws['A4'] = 'Passed:'
        ws['B4'] = quizzes['passed']
        ws['A5'] = 'Failed:'
        ws['B5'] = quizzes['failed']
        ws['A6'] = 'Pass Rate:'
        ws['B6'] = f"{quizzes['pass_rate']:.1f}%"
        ws['A7'] = 'Average Score:'
        ws['B7'] = f"{quizzes['average_score']:.1f}%"
        
        # Adjust column widths
        ws.column_dimensions['A'].width = 20
        ws.column_dimensions['B'].width = 15


def export_compliance_report_pdf(start_date=None, end_date=None) -> HttpResponse:
    """Export compliance report as PDF."""
    generator = ComplianceReportGenerator(start_date, end_date)
    pdf_bytes = generator.generate_pdf_report()
    
    response = HttpResponse(pdf_bytes, content_type='application/pdf')
    response['Content-Disposition'] = f'attachment; filename="compliance_report_{timezone.now().strftime("%Y%m%d")}.pdf"'
    
    return response


def export_compliance_report_excel(start_date=None, end_date=None) -> HttpResponse:
    """Export compliance report as Excel."""
    generator = ComplianceReportGenerator(start_date, end_date)
    excel_bytes = generator.generate_excel_report()
    
    response = HttpResponse(
        excel_bytes,
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    )
    response['Content-Disposition'] = f'attachment; filename="compliance_report_{timezone.now().strftime("%Y%m%d")}.xlsx"'
    
    return response
